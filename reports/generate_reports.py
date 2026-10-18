"""Generate reviewable Vietnamese handover documents from the supplied Jira export.
Usage: python reports/generate_reports.py /path/to/Jira_N1_cap_nhat_18-10-2026.xlsx
Requires openpyxl and python-docx. Does not modify Jira or assert daily work history.
"""
import sys, json, re, hashlib, math
from pathlib import Path
from datetime import datetime, timedelta
from collections import Counter
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from docx import Document
from docx.shared import Inches, Pt
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports'
source = Path(sys.argv[1])
w = load_workbook(source, data_only=True)
raw = list(w['Danh sách Jira'].values)
issues = [dict(zip(raw[0], row)) for row in raw[1:] if row[1]]
byid = {int(i['Issue key'].split('-')[1]): i for i in issues}
fields = ['Issue key','Summary','Assignee','Status','Created','Updated']
evidence = [{k: str(i[k]) if isinstance(i[k],datetime) else i[k] for k in fields} for i in issues]
(OUT/'jira_evidence.json').write_text(json.dumps({'source':source.name,'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'snapshot':'18/10/2026','issues':evidence},ensure_ascii=False,indent=2)+'\n')
people = [
 ('Triệu Quang Dũng','Leader','TRIEU QUANG DUNG'),
 ('Lê Minh Danh','BE','LE MINH DANH'),('Đỗ Tuấn Dũng','BE','DO TUAN DUNG'),
 ('Nguyễn Trần Việt Anh','BE','NGUYEN TRAN VIET ANH'),('Nguyễn Thùy Chang','BE','NGUYEN THUY CHANG'),
 ('Phạm Xuân Bình','BE','PHAM XUAN BINH'),('Hứa Huệ An','FE','HUA HUE AN'),
 ('Hoàng Lan Anh','FE','HOANG LAN ANH'),('Châu Thùy Dung','FE','CHAU THUY DUNG'),('Nguyễn Ngọc Anh','FE','Ngoc Anh')]
# Use actual assignees from the export, not a guessed assignment by stack.
def assigned(person, lo, hi):
 return sorted([i for n,i in byid.items() if lo<=n<=hi and (i['Assignee'] or '').casefold()==person[2].casefold()],key=lambda i:int(i['Issue key'].split('-')[1]))
def brief(items):
 return '; '.join(i['Issue key']+' — '+i['Summary'].strip()+' ['+i['Status']+']' for i in items) or 'Không có issue được giao đích danh trong phạm vi này.'
def keys(items): return ', '.join(i['Issue key'] for i in items) or 'Chưa có phân công đích danh'
source_note = ('Nguồn: '+source.name+' / Danh sách Jira, ảnh chụp trạng thái ngày 18/10/2026. '
 'File không chứa lịch sử chuyển trạng thái, nhật ký theo ngày, trường Parent hoặc lịch Sprint. '
 'Created/Updated là thời điểm tạo/cập nhật issue, không chứng minh ngày thực hiện hay hoàn tất. '
 'Các nhóm S3/S4 của subtasks được đối chiếu theo nội dung, không phải quan hệ Parent lấy từ Jira. '
 'Nhật ký là DỰ THẢO CẦN XÁC NHẬN, không phải bản ghi công việc đã xác thực.')

# API contract: curated workflows plus a mechanically checked route inventory.
sections = [
 ('Xác thực và thư thật','N1-3, N1-4; đối chiếu mã hiện hành',[
 ('POST','/api/auth/login',{'email':'user@example.test','password':'YourPassword9!'},'200: token và thông tin người dùng; 401 nếu sai thông tin.'),
 ('POST','/api/auth/forgot-password',{'email':'user@example.test'},'Yêu cầu email khôi phục; không đồng nhất SMTP chấp nhận với thư đã đến hộp thư.'),
 ('POST','/api/auth/send-register-otp',{'email':'new@example.test'},'OTP email; giới hạn gửi lại 60 giây. Email đã có không được đăng ký mới.'),
 ('POST','/api/auth/register',{'full_name':'Người dùng mới','email':'new@example.test','otp_code':'123456','password':'YourPassword9!','confirm_password':'YourPassword9!'},'Vai trò cố định Sales Rep; xác minh OTP và mật khẩu theo cấu hình backend. Mã OTP dùng trường otp_code; xác nhận mật khẩu dùng confirm_password.'),
 ]),
 ('Khách hàng: CRUD và phân quyền','S3-01 / N1-187; N1-196,197,198',[
 ('GET','/api/customers?search=Minh&status=Tiềm%20năng',None,'200: {success,scope,total,data}; bộ lọc industry, company_size, region, owner_id, custom_fields.'),
 ('POST','/api/customers',{'company_name':'Công ty Minh An','tax_code':'0123456789','email':'sales@example.test','phone':'0912345678','status':'Tiềm năng'},'201: id. Tên bắt buộc; MST nếu nhập phải hợp lệ và duy nhất; owner phải trong phạm vi cho phép.'),
 ('PUT','/api/customers/:id',{'company_name':'Công ty Minh An','parent_id':None},'200: cập nhật hồ sơ được quyền sửa; không tạo chu kỳ cây công ty.'),
 ('DELETE','/api/customers/:id',None,'Kiểm tra quyền và dữ liệu liên quan; không xóa cưỡng bức dữ liệu đang được tham chiếu.'),
 ]),
 ('Người liên hệ','S3-02 / N1-188; N1-199..202',[
 ('GET','/api/contacts/customer/:customerId',None,'Danh sách liên hệ của khách hàng trong phạm vi được phép.'),
 ('POST','/api/contacts/customer/:customerId',{'full_name':'Nguyễn An','email':'an@example.test','phone':'0912345678','decision_role':'Người dùng cuối','is_primary':True},'201: contactId; kiểm tra khách hàng và vai trò mua hàng hợp lệ.'),
 ('PUT','/api/contacts/:id/transfer',{'targetCustomerId':2,'transferReason':'Chuyển công tác'},'Chuyển liên hệ, kiểm tra quyền cả hai hồ sơ và ghi lịch sử.'),
 ('PUT','/api/contacts/:id',{'full_name':'Nguyễn An','is_primary':True},'Cập nhật và bảo đảm quy tắc liên hệ chính.'),
 ]),
 ('Hồ sơ 360 và công ty mẹ/con','S3-03,05 / N1-189,191; N1-203..205,209,210',[
 ('GET','/api/customers/:id/360',None,'Tổng hợp hồ sơ, liên hệ, cơ hội, hoạt động, tệp và cây mẹ/con; dữ liệu và doanh thu tuân thủ scope.'),
 ('PUT','/api/customers/:id',{'parent_id':1},'Liên kết công ty mẹ; không tự liên kết hoặc tạo vòng. Chỉ định ID thực tế được phép truy cập.'),
 ]),
 ('Trùng và gộp khách hàng','S3-04 / N1-190; N1-206..208',[
 ('GET','/api/customers/detect-duplicates',None,'Gợi ý cặp trùng theo chuẩn hóa tên/MST/website. Đây là gợi ý so sánh, khác quy tắc chặn dòng import giống hoàn toàn.'),
 ('POST','/api/customers/merge',{'target_id':1,'source_id':2},'Giữ target, chuyển dữ liệu liên quan từ source trong transaction; cần quyền quản lý và quyền sửa cả hai hồ sơ. Không gộp mẹ vào con.'),
 ]),
 ('Excel khách hàng và lead','S3-06 / N1-192,211..213; S4-02 / N1-232,244..248',[
 ('GET','/api/customers/import/template',None,'XLSX binary: Dữ liệu, Hướng dẫn, Danh mục. Lead thay customers bằng leads.'),
 ('POST','/api/customers/import/preview',{'rows':[{'company_name':'Công ty nhập','email':'import@example.test'}]},'200: {data:[{row,valid,duplicate,duplicate_blocked,error,data}],valid,invalid}; hoặc multipart field file. Xem trước rollback toàn bộ, không ghi dữ liệu thật.'),
 ('POST','/api/customers/import/execute',{'rows':[{'company_name':'Công ty nhập','email':'import@example.test'}]},'200: {data:[{row,success,id|message,duplicate_blocked}],imported,skipped}; mỗi dòng transaction riêng, lỗi một dòng không hủy dòng hợp lệ khác.'),
 ('GET','/api/leads/export',None,'XLSX mẫu tiếng Việt. File xuất nhập lại phải nhận diện dòng đã có, không tạo thêm bản sao.'),
 ('POST','/api/leads/import/preview',{'rows':[{'full_name':'Người liên hệ','company_name':'Công ty nhập','email':'lead@example.test','source_id':1}]},'Giống cơ chế customers; nguồn lead bắt buộc. Trùng hoàn toàn trong DB hoặc dòng trước của tệp bị đánh dấu không hợp lệ.'),
 ('POST','/api/leads/import/execute',{'rows':[{'full_name':'Người liên hệ','company_name':'Công ty nhập','email':'lead@example.test','source_id':1}]},'Thực thi kiểm tra lại ở backend, kể cả gọi trực tiếp không qua preview. Không tin cờ valid do client gửi.'),
 ]),
 ('Tìm kiếm, CSKH và chăm sóc','S3-07..09 / N1-193..195; N1-214..221',[
 ('GET','/api/customers/care-list?days=30',None,'Khách quá N ngày chưa được liên hệ; ưu tiên doanh thu hợp đồng; giới hạn days 1..3650.'),
 ('PUT','/api/customers/:id/contacted',{},'Ghi nhận đã liên hệ và cập nhật thời gian chăm sóc.'),
 ('POST','/api/customers/tickets',{'customer_id':1,'title':'Yêu cầu hỗ trợ','priority':'Cao','status':'Mới'},'Tạo ticket theo schema endpoint; từ 2 ticket chưa giải quyết hiển thị nguy cấp. N1-216,217 chưa có assignee.'),
 ('GET','/api/customers/saved-filters',None,'Bộ lọc của người dùng; POST để lưu, DELETE /:id để xóa. UI có thể lưu cục bộ theo luồng hiện hành.'),
 ]),
 ('Form công khai, lead và campaign','S4-01..03 / N1-231..233; N1-240..252',[
 ('POST','/api/leads/forms',{'name':'Form website','source_id':1},'Quản lý tạo form; key sinh bởi hệ thống. GET /api/public/lead-forms/:key/meta để lấy cấu hình.'),
 ('POST','/api/public/lead-forms/:key',{'full_name':'Khách mới','company_name':'Công ty mới','email':'new-lead@example.test'},'Không Bearer; kiểm tra form, nguồn mặc định và chống spam/IP. Không đưa token quản trị vào mã nhúng.'),
 ('POST','/api/leads',{'full_name':'Lead mới','company_name':'Doanh nghiệp mới','email':'lead@example.test','source_id':1},'201: {id}; cần email hoặc điện thoại. Không tự đặt score/temperature bằng payload.'),
 ('GET','/api/campaigns',None,'Danh sách và thống kê chiến dịch theo quyền; Sales lấy tên bằng /api/campaigns/options.'),
 ('POST','/api/campaigns',{'name':'Chiến dịch Q4','budget':10000000,'channel':'Website','start_date':'2026-10-18','end_date':'2026-11-18'},'Tạo chiến dịch; ngân sách và ngày phải hợp lệ.'),
 ]),
 ('Trùng, chấm điểm và phân bổ lead','S4-04..06 / N1-234..236; N1-253..264',[
 ('GET','/api/leads/duplicates',None,'Gợi ý lead trùng; POST /api/leads/merge với target_id,source_id,reason để gộp theo quyền.'),
 ('GET','/api/leads/scoring-rules',None,'Admin/Director: {data,thresholds:{warm,hot}}.'),
 ('PUT','/api/leads/scoring-rules',{'warm':30,'hot':70,'rules':[{'field_name':'source_id','match_value':'1','points':50}]},'Admin/Director; tối đa 100 quy tắc; field industry/company_size/source_id/interest_level; points -1000..1000, hot>warm. Tính lại lead đang hoạt động trong transaction.'),
 ('PUT','/api/leads/assignment-rules',{'rules':[]},'Đặt danh sách quy tắc; lấy cấu trúc hiện tại bằng GET cùng URL trước khi sửa. Mảng rỗng xóa quy tắc, chỉ dùng có chủ đích trên DB test.'),
 ('POST','/api/leads/:id/assign',{'owner_id':4,'reason':'Phân theo khu vực','sla_minutes':60},'Nhân viên tiếp nhận phải active và có Sales Rep; kiểm tra scope. Worker phân vòng theo rule/group.'),
 ]),
 ('Phản hồi, chuyển đổi và lọc lead','S4-07..09 / N1-237..239; N1-265..276',[
 ('POST','/api/leads/:id/accept',{},'Chỉ nhân viên đang được phân lead ASSIGNED được nhận; chuyển sang CARING.'),
 ('POST','/api/leads/:id/reject',{'reason':'Khách thuộc khu vực khác'},'Bắt buộc lý do; trả về hàng chờ theo quy trình.'),
 ('POST','/api/leads/:id/convert',{'deal':{'title':'Triển khai CRM','stage_id':1,'amount':5000000,'expected_close_date':'2026-12-01'}},'Một transaction tạo/chọn customer, tạo contact/deal, chuyển lịch sử và đánh dấu lead CONVERTED; rollback nếu bất kỳ bước nào lỗi.'),
 ('GET','/api/leads?temperature=HOT&source_id=1',None,'Lọc theo source/campaign/owner/status/temperature/SLA và ngày theo leadRoutes.list; dữ liệu có scope, không tự loại lead điểm thấp.'),
 ]),
 ('Cây tổ chức và báo cáo cơ hội','Mã hiện hành; cây tổ chức liên quan S2, Pipeline không có Sprint 5 trong file Jira mới',[
 ('GET','/api/organization/tree',None,'Một cây phân cấp: children, leader, region, member_count trực tiếp; UI tổng hợp total_members của nhánh.'),
 ('POST','/api/organization/groups',{'name':'Nhóm kinh doanh','leader_id':3,'parent_id':None,'region_id':None},'Quyền CATEGORY write; kiểm tra trưởng nhóm/nhóm cha; không tạo chu kỳ.'),
 ('GET','/api/groups/:id/members?include_descendants=1',None,'Nhân sự toàn nhánh để đối chiếu quy mô; không cộng trùng một người trong cùng danh sách.'),
 ('GET','/api/deals/export?status=OPEN',None,'XLSX tiếng Việt, scope và bộ lọc giống danh sách; tên liên hệ/owner/nguồn/chiến dịch/giai đoạn, số tiền numeric, ngày ký/dự kiến, lý do thắng/thua.'),
 ])]

md=['# API N1 — bàn giao theo Jira','',source_note,'',
 '## Quy ước tích hợp','Base path `/api`. Gửi `Authorization: Bearer <token>` cho route nội bộ, `Content-Type: application/json` khi gửi JSON. Các ID ví dụ là minh họa; lấy ID thật từ DB test/danh mục. Không gửi thư cho email ví dụ `.test`.',
 'Scope và quyền đọc/ghi lấy từ role matrix; Admin toàn quyền. Các vai trò khác cần kiểm tra cả module và phạm vi OWN/TEAM/ALL. Ẩn nút ở frontend không thay kiểm tra backend.',
 'Lỗi chung: `{success:false,message,details?}`; 400 dữ liệu không hợp lệ; 401 thiếu/sai token; 403 thiếu quyền; 404 không tìm thấy/không thuộc phạm vi; 409 xung đột; 413 tệp quá lớn; 422 điều kiện nghiệp vụ; 429 giới hạn yêu cầu; 500 lỗi máy chủ. Endpoint có thể dùng tập mã con; xem test và route tương ứng.',
 'Danh sách thường có `{data,total?,scope?}`; POST create thường trả 201 và id. Một số API cũ dùng contactId hoặc cấu trúc riêng; frontend phải theo endpoint, không giả định toàn bộ response giống nhau.',
 'XLSX: client tải bằng arraybuffer/blob, không parse JSON. Import dùng multipart `file` hoặc JSON `{rows:[...]}`, tối đa 1.000 dòng, tệp 5 MB. Dòng trùng hoàn toàn không có lựa chọn ép nhập/cập nhật; cùng MST nhưng thay thông tin có thể chọn update theo quyền và xem trước lại.',
 'Điểm lead tính tự động theo AC S4-05; thay quy tắc/ngưỡng bằng quyền Admin/Director. 0 điểm là hợp lệ khi chưa có tiêu chí khớp. Không thêm ô ghi đè score trái AC.']
for title, jira, endpoints in sections:
 md += ['', '## '+title, 'Đối chiếu: '+jira+'.']
 for meth,path,body,note in endpoints:
  md += ['', '### '+meth+' '+path, note]
  if body is not None: md += ['```json',json.dumps(body,ensure_ascii=False,indent=2),'```']
md+=['','## Trách nhiệm theo thành viên','Danh sách BE/FE theo ảnh nhóm; task cụ thể theo Assignee của Jira. Việc FE/BE hỗ trợ nhau không được đổi assignee trong báo cáo.']
for person in people:
 md += ['', '### '+person[0]+' ('+person[1]+')','Sprint 3: '+brief(assigned(person,196,221)),'Carry-over import user: '+brief(assigned(person,223,230)),'Sprint 4 (tất cả To Do): '+brief(assigned(person,240,276))]
md+=['','## Điểm cần thống nhất trước khi báo cáo',
 'S3: 3/9 story Done; 6/9 In Progress dù subtasks ngày 10/10 đều Done. Không suy diễn story đã nghiệm thu chỉ từ subtask.',
 'S4: 46 issue mới (9 story + 37 subtasks) đều To Do tại 18/10. Tài liệu API mô tả mã hiện có, không phải chứng nhận các issue Jira này Done.',
 'N1-216 và N1-217 chưa có assignee; không gán cho Phạm Xuân Bình. Jira ghi Nguyễn Thùy Chang, khác tên Trang ở ảnh mẫu nhật ký. Giữ tên Jira và chờ nhóm xác nhận.',
 'N1-222 là carry-over nhập người dùng, không có nhãn Sprint rõ trong dữ liệu export. N1-229 To Do ngày 18/10; không ghi đã kiểm thử xong.',
 'SLA: task N1-266/268 nêu 3 ngày, API hiện có sla_minutes (mặc định 60 phút). Nhóm cần chốt cấu hình nghiệm thu 4.320 phút nếu giữ AC 3 ngày; không đồng nhất cấu hình mặc định với yêu cầu Jira.',
 'Không có Sprint 5 trong file đính kèm. Các endpoint Pipeline ghi theo mã hiện hành, không gắn mã Jira tự tạo.',
 'Gửi email: kiểm tra SMTP runtime sau restart backend; log queue/attempt và thư hộp đến là bằng chứng khác nhau. Không đưa mật khẩu SMTP, OTP thật hoặc token reset vào tài liệu.']
# Route inventory generated from actual source. Dynamic routes explicitly added.
md+=['','## Phụ lục route đọc trực tiếp từ source','Chỉ liệt kê route khai báo tĩnh; không coi bảng này là mô tả đầy đủ schema/quyền từng route. Route động accept/reject và import được bổ sung phía dưới.']
for name,prefix in [('auth','auth'),('user','users'),('customer','customers'),('contact','contacts'),('lead','leads'),('campaign','campaigns'),('organization','organization'),('group','groups'),('deal','deals'),('quote','quotes')]:
 text=(ROOT/f'backend/src/routes/{name}Routes.js').read_text()
 routes=re.findall(r'(router|r|publicRouter)\.(get|post|put|patch|delete)\(\s*["\']([^"\']+)',text)
 md+=['','### '+prefix]
 for variable,method,path in routes:
  if name=='lead' and path=='/:id/': continue
  base='public' if variable=='publicRouter' else prefix
  md += ['- `'+method.upper()+' /api/'+base+('' if path=='/' else path)+'`']
md+=['','Routes động: `POST /api/leads/:id/accept`, `POST /api/leads/:id/reject`. Với customers/leads/users: `GET /import/template`, `POST /import/preview`, `POST /import/execute`, `POST /import-batch`; riêng customers thêm `/bulk-import`. Customers/leads có GET `/export` tại importRoutes.']
(OUT/'API_N1_Theo_Jira.md').write_text('\n'.join(md)+'\n')

def doc_from_md(lines, dest):
 d=Document();sec=d.sections[0];sec.header.paragraphs[0].text='ENTERPRISE CRM · NHÓM N1 · 18/10/2026'
 sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7)
 style=d.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(10)
 code=False
 for line in lines:
  if line.startswith('```'): code=not code;continue
  if line.startswith('#'):
   level=len(line)-len(line.lstrip('#'));d.add_heading(line[level:].strip(),min(level,3))
  elif line:
   p=d.add_paragraph(line)
   if code:
    for run in p.runs:run.font.name='Consolas';run.font.size=Pt(9)
 sec.footer.paragraphs[0].text='Nguồn Jira và trạng thái là căn cứ; phần đề xuất/dự thảo cần nhóm xác nhận.'
 d.save(dest)
doc_from_md(md,OUT/'API_N1_Theo_Jira.docx')

review=['# Đề xuất Sprint Review','',source_note,'',
 '## Lựa chọn đề xuất: Sprint 3, sau khi chốt nghiệm thu',
 'Sprint 3 cho câu chuyện demo liền mạch và sát nghiệp vụ nhất: tạo doanh nghiệp → thêm người liên hệ → hồ sơ 360 → công ty mẹ/con → kiểm tra và gộp trùng → nhập Excel → lọc và chăm sóc. Đây là đề xuất về giá trị demo; chưa phải kết luận toàn bộ Sprint 3 Done.',
 '## So sánh có căn cứ',
 'Sprint 2: 10 story chính trong nhóm S2 của file đang Done (N1-115,125,126,127,118,119,120,121,122,123). Phù hợp nhất nếu phải review ngay dựa trên trạng thái Done hiện tại; cần rà lại tên S2-02/N1-125 đang trùng chủ đề liên hệ với S3-02.',
 'Sprint 3: N1-188,189,193 Done; N1-187,190,191,192,194,195 In Progress. Có thể chọn làm buổi review có giá trị cao nếu nghiệm thu được các story còn lại; không tự đổi trạng thái từ tài liệu.',
 'Sprint 4: N1-231..276, 46 issue đều To Do. Chưa đủ bằng chứng để trình bày là sprint đã hoàn tất. Sprint 5 không có trong file.',
 '## Kịch bản 15 phút cho Sprint 3',
 'Phút 0–2: nêu mục tiêu, phạm vi và bảng trạng thái story. Tạo khách hàng và đổi vai trò để chứng minh quyền OWN/TEAM/ALL (N1-187,196..198).',
 'Phút 2–4: thêm liên hệ chính, vai trò mua hàng và chuyển công ty có lịch sử (N1-188,199..202).',
 'Phút 4–6: mở 360, đối chiếu KPI cơ hội/doanh thu Won và dữ liệu nhánh công ty mẹ/con; chỉ tính dữ liệu được phép xem (N1-189,191,203..205,209,210).',
 'Phút 6–9: so sánh/gộp hai hồ sơ trùng và chứng minh liên hệ, cơ hội, hoạt động, tệp còn nguyên; chặn trường hợp gộp gây vòng mẹ/con (N1-190,206..208).',
 'Phút 9–11: tải mẫu Excel, nhập tệp có dòng hợp lệ, trùng hoàn toàn và sai định dạng; kiểm tra preview không ghi DB và execute không nhập dòng trùng (N1-192,211..213).',
 'Phút 11–13: tìm kiếm/lưu bộ lọc; hiển thị khách nguy cấp từ ticket chưa xử lý và danh sách quá N ngày chưa chăm sóc (N1-193..195,214..221). Không nói ticket đã phân công cho cá nhân khi Jira chưa ghi.',
 'Phút 13–15: đối chiếu test, xin phản hồi theo từng AC; ghi lỗi còn mở và quyết định chấp nhận/tồn đọng. Chỉ cập nhật Jira khi nhóm thực sự nghiệm thu.',
 '## Điều kiện trước buổi review',
 'Đối chiếu AC-MATRIX.md với yêu cầu gốc của nhóm; chạy DB test riêng, không dùng dữ liệu khách thật; chuẩn bị seed theo vai trò, lưu bằng chứng request/response và ảnh trước/sau. Chốt parent story In Progress hoặc công khai chưa hoàn thành. Những lỗi sửa sau snapshot 18/10 ghi riêng là sửa bổ sung, không giả ngày hoàn tất vào Jira.',
 '## Câu chốt cho báo cáo',
 '“Nhóm đề xuất Sprint 3 làm nội dung review trọng tâm do thể hiện trọn luồng quản lý khách hàng. Tại bản Jira ngày 18/10, Sprint 3 còn 6 story In Progress; nội dung nghiệm thu chỉ được xác nhận sau khi kiểm tra AC. Nếu yêu cầu chọn sprint đã Done theo dữ liệu hiện có, chọn Sprint 2.”']
(OUT/'De_xuat_Sprint_Review.md').write_text('\n'.join(review)+'\n');doc_from_md(review,OUT/'De_xuat_Sprint_Review.docx')

# Journals: no invented actual activity dates. Every daily narrative is labelled draft.
b=Workbook();b.remove(b.active)
headers=['Thành viên','HÔM QUA LÀM GÌ','GẶP KHÓ KHĂN GÌ KHÔNG','HÔM NAY LÀM GÌ']
blue='173B63';line=Side(style='thin',color='DCE5EF')
def sheet(name, rows, widths, note=None):
 s=b.create_sheet(name);s.append(rows[0])
 for row in rows[1:]:s.append(row)
 s.freeze_panes='B2';s.auto_filter.ref=s.dimensions
 for j,width in enumerate(widths,1):s.column_dimensions[__import__('openpyxl').utils.get_column_letter(j)].width=width
 for row in s:
  s.row_dimensions[row[0].row].height=34 if row[0].row==1 else min(409, max(34, 15 * max(sum(max(1,math.ceil(len(part)/max(10,widths[j]-4))) for part in str(cell.value or '').split('\n')) for j,cell in enumerate(row)) + 12))
  for cell in row:
   cell.alignment=Alignment(vertical='top',wrap_text=True)
   cell.border=Border(bottom=line,right=line)
   if cell.row==1:cell.fill=PatternFill('solid',fgColor=blue);cell.font=Font(bold=True,color='FFFFFF')
   elif cell.row%2==0:cell.fill=PatternFill('solid',fgColor='F1F6FC')
 s.print_title_rows='1:1';s.sheet_properties.pageSetUpPr.fitToPage=True
 s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
 return s
intro=[['Mục','Nội dung'],['Tính chất','BẢN DỰ THẢO CẦN XÁC NHẬN. Không dùng như bằng chứng nhật ký thực tế nếu chưa đối chiếu thành viên.'],['Nguồn',source_note],['Phạm vi ngày','07–18/10/2026 là cửa sổ có dữ liệu S3 đến snapshot, không phải lịch Sprint được xác nhận. Cả ngày cuối tuần đều là bản nháp theo lịch, không khẳng định có đi làm.'],['Cách dùng','Giữ 4 cột như ảnh mẫu. Chọn sheet theo ngày, xác nhận hôm qua/hôm nay và thay phần khó khăn bằng sự việc thật trước khi đưa lên internship. Sheet Cá nhân Leader và Tổng hợp nhóm thuận tiện sao chép.'],['Khó khăn','Jira không ghi khó khăn từng ngày. Ghi “chưa có dữ liệu, cần xác nhận”, không tự đặt SMTP lỗi hay thời gian hoàn thành cho cá nhân.'],['Sprint 4','Sheet Kế hoạch S4 chứa công việc To Do theo đúng assignee; đây là kế hoạch tiếp theo, không phải nhật ký hoàn thành.'],['Sprint 5 trở đi','Không có trong file. Cần bổ sung Jira để lập nhật ký có căn cứ; không tạo issue hoặc ngày làm việc giả.'],['Tên thành viên','Dùng Nguyễn Thùy Chang theo Jira; ảnh mẫu có chữ Trang. Nguyễn Ngọc Anh khớp assignee Ngoc Anh.'],['Quan hệ task','Nhóm story/subtask theo nội dung, không khẳng định parent link vì export không có Parent. N1-222..230 là carry-over chưa có Sprint rõ.'],['Dữ liệu đối chiếu','Sheet Jira nguồn lưu key, summary, assignee, trạng thái, Created, Updated; không xuất các ID tài khoản Atlassian.']]
sheet('Đọc trước',intro,[28,120])
phases={7:'Đề xuất đầu kỳ: rà yêu cầu, đầu vào/đầu ra và tiêu chí chấp nhận; tách phần kiểm tra theo nhiệm vụ được giao.',8:'Đề xuất triển khai: hoàn thiện logic/giao diện chính, thống nhất dữ liệu mẫu và hợp đồng API với phần phụ thuộc.',9:'Đề xuất kiểm tra: rà dữ liệu rỗng/sai định dạng, phân quyền và tình huống biên; ghi câu hỏi cần thống nhất.',10:'Đối chiếu snapshot: các subtask S3 dưới đây được ghi Done và Updated 10/10; cần xác nhận hoạt động thực tế trong ngày, không coi Updated là ngày hoàn tất.',11:'Đề xuất rà tích hợp: chạy lại luồng từ danh sách đến chi tiết, kiểm tra kết quả backend/frontend và ghi lỗi còn mở.',12:'Đề xuất nghiệm thu liên hệ/360/mẹ-con: đối chiếu phần phụ thuộc; N1-188 và189 hiện Done, N1-191 còn In Progress.',13:'Đề xuất hồi quy: kiểm tra quyền theo vai trò, dữ liệu liên kết và trường hợp cập nhật/xóa; tổng hợp bằng chứng kiểm thử.',14:'Đề xuất rà bộ lọc và carry-over import user: N1-193 Done, N1-222..230 vẫn To Do; thống nhất đầu việc chưa hoàn tất.',15:'Đề xuất củng cố kiểm thử: rà các story S3 còn In Progress và tiêu chí chưa có bằng chứng chấp nhận.',16:'Đề xuất chuẩn bị review: chọn kịch bản khách hàng/liên hệ/360/import/gộp, đối chiếu issue và dữ liệu demo.',17:'Đề xuất rà carry-over: N1-144 và174 của Leader hiện Done, Updated 17/10; kiểm tra ảnh hưởng đến nghiệm thu S3.',18:'Đề xuất lập kế hoạch S4: 46 issue S4 được tạo và đang To Do; rà trách nhiệm, phụ thuộc và tiêu chí kiểm thử trước khi nhận việc.'}
block='Jira không có ghi chú khó khăn theo ngày; cần thành viên xác nhận. Không đồng nhất trạng thái Done của subtask với nghiệm thu story.'
leader_rows=[['Ngày / Thành viên',*headers[1:]]];group_rows=[['Ngày / Phạm vi',*headers[1:]]]
trace=[['Ngày dự thảo','Thành viên','Jira tham chiếu','Căn cứ','Đã xác nhận?']]
for day in range(7,19):
 date=f'{day:02d}/10/2026';rows=[headers];summary=[]
 for person in people:
  s3=assigned(person,196,221);carry=assigned(person,223,230);s4=assigned(person,240,276)
  base=s4 if day==18 else carry if day==14 and carry else s3
  today='DỰ THẢO CẦN XÁC NHẬN — '+phases[day]+'\nPhạm vi cá nhân: '+brief(base)
  if person[1]=='Leader':
   today+='\nĐề xuất điều phối: tổng hợp vấn đề tích hợp, rà quyền và đối chiếu trạng thái story; hoạt động điều phối chưa được ghi theo ngày trong Jira.'
   if day==17:today+='\nCarry-over S2: '+brief([byid[144],byid[174]])
  prev=day-1
  yesterday=('Chưa có dữ liệu trước cửa sổ S3 ngày 07/10; cần xác nhận công việc hôm trước.' if day==7 else 'DỰ THẢO CẦN XÁC NHẬN — tiếp nối phần '+keys(assigned(person,223,230) if prev==14 and carry else s3)+'. '+phases[prev])
  rows.append([person[0]+' ('+person[1]+')',yesterday,block,today])
  trace.append([date,person[0],keys(base),'Phân bổ hoạt động theo ngày là đề xuất; status lấy snapshot 18/10.','Chưa'])
  summary.append(person[0]+': '+keys(base))
  if person[1]=='Leader':leader_rows.append([date+' · '+person[0],yesterday,block,today])
 sheet(f'{day:02d}-10 Dự thảo',rows,[28,65,48,95])
 group_rows.append([date+' · Cả nhóm',
  'DỰ THẢO CẦN XÁC NHẬN — '+('Chưa có nhật ký hôm trước trong nguồn.' if day==7 else phases[day-1]),
  'Chưa có blocker thực tế theo ngày trong Jira. Cần xử lý lệch trạng thái story/subtask S3 và phân công N1-216,217 còn trống; không gán các vấn đề này là khó khăn đã xảy ra của một thành viên.',
  'DỰ THẢO CẦN XÁC NHẬN — '+phases[day]+'\n'+ '\n'.join(summary)+'\nN1-216,217: chưa phân công; cần Leader thống nhất assignee.'])
sheet('Cá nhân Leader',leader_rows,[32,65,48,95]);sheet('Tổng hợp nhóm',group_rows,[28,65,55,110])
plans=[['Thành viên','HÔM QUA LÀM GÌ','GẶP KHÓ KHĂN GÌ KHÔNG','HÔM NAY LÀM GÌ']]
for person in people:
 plans.append([person[0]+' ('+person[1]+')','Chưa có Daily Scrum S4 trong file. Không ghi công việc To Do là đã làm.','Cần chốt hợp đồng API, dữ liệu danh mục và SLA 3 ngày/60 phút trước nghiệm thu; đây là điểm cần thống nhất, chưa phải blocker có chứng cứ.','KẾ HOẠCH S4 — CHƯA THỰC HIỆN THEO JIRA\n'+brief(assigned(person,240,276))])
sheet('Kế hoạch S4',plans,[28,60,65,110])
sheet('Jira nguồn',[fields]+[[i[k].isoformat(' ') if isinstance(i[k],datetime) else i[k] for k in fields] for i in issues],[18,100,28,20,25,25])
sheet('Đối chiếu nhật ký',trace,[22,30,65,85,20])
b.save(OUT/'Nhat_ky_Sprint_3_tro_di.xlsx')
# A textual personal/team journal is easier to paste into internship than spreadsheet alone.
logmd=['# Nhật ký cá nhân và tổng hợp nhóm từ Sprint 3','',source_note,'','Các nội dung theo ngày dưới đây là dự thảo đề xuất. Trạng thái issue được giữ đúng ảnh chụp Jira, không ghi ngày Done suy từ Updated.']
for title,rows in [('Cá nhân Triệu Quang Dũng',leader_rows),('Tổng hợp nhóm',group_rows)]:
 logmd += ['','## '+title]
 for row in rows[1:]:
  logmd += ['','### '+row[0],'HÔM QUA LÀM GÌ: '+row[1],'GẶP KHÓ KHĂN GÌ KHÔNG: '+row[2],'HÔM NAY LÀM GÌ: '+row[3]]
(OUT/'Nhat_ky_ca_nhan_va_nhom.md').write_text('\n'.join(logmd)+'\n');doc_from_md(logmd,OUT/'Nhat_ky_ca_nhan_va_nhom.docx')
print(json.dumps({'issues':len(issues),'states':dict(Counter(i['Status'] for i in issues)),'journal_sheets':len(b.sheetnames),'member_day_drafts':120,'output':str(OUT)},ensure_ascii=False))
